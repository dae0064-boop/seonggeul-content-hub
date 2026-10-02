#!/usr/bin/env python3
"""
스레드 보험 심의글 묶음을 한 편씩 Word(.docx) 로 만든다. 본문 밑에 고정 사진(심의 고지 배너)을 붙인다.

  python3 scripts/threads/insurance-word.py <묶음.md> <사진.jpg> <저장 폴더> [--skip-fresh]

--skip-fresh: 이미 있는 Word 가 묶음 파일·사진보다 새것이면 건너뛴다 (PC 실행기가 쓴다).
  사람이 Drive 에서 고친 Word 를 덮어쓰지 않고, 원고나 사진이 바뀌었을 때만 다시 만든다.

결과: <저장 폴더>/<날짜>_<번호>_<글감 앞부분>.docx  (5개)

사진에는 설계사 이름·등록번호가 들어 있다. 사진과 만든 Word 파일은 저장소에 커밋하지 않는다
(저장소가 공개다). Google Drive ClaudeWorkspace/보험글/<날짜>/ 에 둔다.
필요한 것: pip install python-docx
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor

FONT = '맑은 고딕'


def parse_set(text):
    text = text.replace('\r\n', '\n')
    meta = {}
    fm = re.match(r'^---\n(.*?)\n---\n', text, re.S)
    if fm:
        for line in fm.group(1).split('\n'):
            if ':' in line:
                k, v = line.split(':', 1)
                meta[k.strip()] = v.strip()
        text = text[fm.end():]
    posts = []
    parts = re.split(r'^##\s+(\d+)\s*$', text, flags=re.M)
    for i in range(1, len(parts), 2):
        lines = parts[i + 1].lstrip('\n').split('\n')
        fields, j = {}, 0
        while j < len(lines):
            m = re.match(r'^(역할|보험|글감|출처|원문확인):\s*(.*)$', lines[j])
            if not m:
                break
            fields[m.group(1)] = m.group(2).strip()
            j += 1
        fields['no'] = int(parts[i])
        fields['body'] = '\n'.join(lines[j:]).strip()
        posts.append(fields)
    return meta, posts


def set_font(run, size, color=None, bold=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', FONT)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def para(doc, text, size=11, color=None, bold=False, after=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.4
    set_font(p.add_run(text), size, color, bold)
    return p


def out_name(post, meta):
    short = re.sub(r'[\\/:*?"<>|?()\'‘’—,·]', '', post.get('글감', ''))
    short = re.sub(r'\s+', ' ', short).strip()[:20].strip()
    return f"{meta.get('date', 'set')}_{post['no']}_{short}.docx"


def build(post, meta, photo, out_dir):
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)

    # 본문 — 스레드에 그대로 붙여 넣을 글. 빈 줄은 빈 문단으로 둔다
    for line in post['body'].split('\n'):
        para(doc, line)

    para(doc, '')
    doc.add_picture(str(photo), width=sec.page_width - sec.left_margin - sec.right_margin)

    # 작업 메모 — 올리지 않는다. 심의 제출·출처 확인용
    para(doc, '')
    grey = (0x80, 0x80, 0x80)
    para(doc, '── 작업 메모 (올리지 않음) ──', 9, grey, bold=True)
    para(doc, f"{meta.get('date', '')} 묶음 {post['no']}번 · {post.get('역할', '')} · {post.get('보험', '')}", 9, grey)
    para(doc, f"글감: {post.get('글감', '')}", 9, grey)
    para(doc, f"출처: {post.get('출처', '')}", 9, grey)
    para(doc, f"출처 원문 확인: {post.get('원문확인', '')}", 9, grey)
    para(doc, f"글자 수: {len(post['body'])}자", 9, grey)

    path = Path(out_dir) / out_name(post, meta)
    doc.save(path)
    return path


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    skip_fresh = '--skip-fresh' in sys.argv
    if len(args) != 3:
        print(__doc__)
        sys.exit(1)
    src, photo, out_dir = args
    Path(out_dir).mkdir(parents=True, exist_ok=True)
    meta, posts = parse_set(Path(src).read_text(encoding='utf-8'))
    newest_input = max(Path(src).stat().st_mtime, Path(photo).stat().st_mtime)
    for post in posts:
        target = Path(out_dir) / out_name(post, meta)
        if skip_fresh and target.exists() and target.stat().st_mtime >= newest_input:
            print(f'그대로 둠: {target.name}')
            continue
        print(f'만듦: {build(post, meta, photo, out_dir).name}')


if __name__ == '__main__':
    main()
